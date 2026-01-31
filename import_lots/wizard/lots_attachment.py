# -*- coding: utf-8 -*-

from odoo import fields, models, _
from odoo.exceptions import UserError, ValidationError
import base64
import openpyxl
from io import BytesIO


class LotsAttachment(models.TransientModel):
    """ Class for lots wizard """
    _name = 'lot.attachment'
    _description = "Lots Attachment"

    picking_id = fields.Many2one('stock.picking',
                                 string="Selección de acciones",
                                 help="Parent picking")
    product_id = fields.Many2one('product.product',
                                 string="Producto",
                                 help="Current product")
    demanded_quantity = fields.Float(string="Cantidad",
                                     help="Product quantity demanded")
    type = fields.Selection(string="Tipo de lotes",
                            selection=[('lot', 'Lot'), ('serial', 'Serial')],
                            help="Choose a lot/serial")
    move_id = fields.Many2one('stock.move')
    attachment = fields.Binary(string="Subir", attachment=True)
    attachment_name = fields.Char(string="Nombre del archivo adjunto",
                                  help="Attachment file name")

    def action_import_lot(self):
        """ Importing lots """
        current_move_id = self.env['stock.move'].browse(self.move_id.id)
        if self.attachment:
            wb = openpyxl.load_workbook(
                filename=BytesIO(base64.b64decode(self.attachment)),
                read_only=True) if self.attachment else ""
            ws = wb.active
            # Check if product exists in the sheet
            product_found = any(
                record[1] == current_move_id.product_id.display_name for
                record in ws.iter_rows(min_row=2, values_only=True))
            if not product_found:
                raise UserError(
                    _('El producto "%s" no existe en la hoja.') %
                    current_move_id.product_id.display_name)
            # Check if lot name already exists in move line ids
            if (current_move_id.move_line_ids and
                    any(record[0] in set(current_move_id.move_line_ids.mapped('lot_name')) for
                    record in ws.iter_rows(min_row=2, values_only=True))):
                raise UserError(
                    _('Este nombre de lote ya existe en la línea de movimiento.'))
            # Calculate total sheet quantity for the product
            total_sheet_quantity = sum(
                record[2] for record in ws.iter_rows
                (min_row=2, values_only=True) if
                record[1] == current_move_id.product_id.display_name)
            # Check if total sheet quantity exceeds demand quantity of the
            # product
            if total_sheet_quantity > current_move_id.product_uom_qty:
                raise UserError(
                    _('Total quantity in the sheet exceeds the demand '
                      'quantity of the product. Please adjust the quantities '
                      'in the sheet.'))
            # Prepare move line values to be written
            vals_list = []
            for record in ws.iter_rows(min_row=2, values_only=True):
                lot_name, product_name, quantity = record
                if product_name == current_move_id.product_id.display_name:
                    if current_move_id.picking_type_id.code == 'outgoing':
                        lote= self.env['stock.lot'].sudo().search([('name','=',lot_name)])
                        if not lote:
                            raise UserError(_('El Lote/Serie "%s" no existe en el sistema.') % lot_name)
                        lote_stock= self.env['stock.quant'].search([
                                                ('lot_id', '=', lote.id),
                                                ('location_id', '=', current_move_id.location_id.id),
                                                ('quantity', '>', 0),
                                            ], limit=1)
                        if not lote_stock:
                            raise UserError(_('El Lote/Serie "%s" No tiene stock en esta ubicación') % lot_name)
                        vals_list.append((0, 0, {
                            'lot_id': lote.id,
                            'quantity': min(quantity, current_move_id.product_qty),
                            'move_id': current_move_id.id,
                        }))
                    else:
                        vals_list.append((0, 0, {
                            'lot_name': lot_name,
                            'quantity': min(quantity, current_move_id.product_qty),
                            'move_id': current_move_id.id,
                        }))
            # Write move line values
            current_move_id.move_line_ids.unlink()
            current_move_id.write({'move_line_ids': vals_list})
            return {
                'type': 'ir.actions.client',
                'tag': 'reload',
            }
        else:
            raise ValidationError(
                _('Check whether you upload the document'))

    def action_download_sample(self):
        """ For downloading a sample excel file """
        return {
            'type': 'ir.actions.act_url',
            'url': '/download/excel',
            'target': 'self',
            'file_name': 'my_excel_file.xlsx'
        }
